"""Build the three-page adviser briefing using the bundled python-docx runtime.

Content and layout source for strategic-shareholder-briefing.docx. Render with the
documents skill renderer; the companion Markdown record holds claim provenance.
"""

from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "strategic-shareholder-briefing.docx"
doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.27)
section.page_height = Inches(11.69)
section.top_margin = Inches(0.70)
section.bottom_margin = Inches(0.66)
section.left_margin = Inches(0.76)
section.right_margin = Inches(0.76)
section.header_distance = Inches(0.28)
section.footer_distance = Inches(0.28)

for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2"):
    style = doc.styles[name]
    style.font.name = "Arial"
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    style.element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
normal = doc.styles["Normal"]
normal.font.size = Pt(10.5)
normal.paragraph_format.line_spacing = 1.12
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.widow_control = True
doc.styles["Title"].font.size = Pt(24)
doc.styles["Title"].font.bold = True
doc.styles["Title"].paragraph_format.space_after = Pt(10)
doc.styles["Heading 1"].font.size = Pt(20)
doc.styles["Heading 1"].paragraph_format.space_after = Pt(12)
doc.styles["Heading 2"].font.size = Pt(12)
doc.styles["Heading 2"].paragraph_format.space_before = Pt(10)
doc.styles["Heading 2"].paragraph_format.space_after = Pt(5)
for name in ("Title", "Subtitle", "Heading 1", "Heading 2"):
    doc.styles[name].paragraph_format.keep_with_next = True
    ppr = doc.styles[name].element.find(qn("w:pPr"))
    if ppr is not None:
        borders = ppr.find(qn("w:pBdr"))
        if borders is not None:
            ppr.remove(borders)

doc.core_properties.title = "Circularo Strategic Shareholder Opportunity"
doc.core_properties.subject = "Confidential briefing for Joshua Jahani"
doc.core_properties.author = "Josef Neumann"
doc.core_properties.keywords = "Circularo, strategic shareholder, confidential, draft"
doc.core_properties.comments = "Draft for Josef Neumann's review. See companion record for provenance."


def paragraph(text, *, size=None, bold=False, after=None):
    p = doc.add_paragraph()
    # Completion fields remain searchable and editable, with a distinct amber tone.
    for part in re.split(r"(\[COMPLETE[^\]]*\])", text):
        if not part:
            continue
        r = p.add_run(part)
        r.bold = bold or part.startswith("[COMPLETE")
        if size:
            r.font.size = Pt(size)
        if part.startswith("[COMPLETE"):
            r.font.color.rgb = RGBColor.from_string("8A4B08")
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    return p


def heading(text):
    doc.add_paragraph(text, "Heading 2")


def table(headers, rows, widths, *, numeric=False):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col, width in zip(t.columns, widths):
        col.width = Inches(width)
    tblpr = t._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        item = OxmlElement("w:" + edge)
        item.set(qn("w:val"), "single")
        item.set(qn("w:sz"), "4")
        item.set(qn("w:color"), "D9D9D9")
        borders.append(item)
    tblpr.append(borders)
    marg = OxmlElement("w:tblCellMar")
    for edge, value in (("top", "100"), ("bottom", "100"), ("left", "115"), ("right", "115")):
        item = OxmlElement("w:" + edge)
        item.set(qn("w:w"), value)
        item.set(qn("w:type"), "dxa")
        marg.append(item)
    tblpr.append(marg)
    repeat = OxmlElement("w:tblHeader")
    t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for vals in rows:
        t.add_row()
    for row_index, vals in enumerate([headers, *rows]):
        row = t.rows[row_index]
        no_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(no_split)
        for col_index, value in enumerate(vals):
            cell = row.cells[col_index]
            cell.width = Inches(widths[col_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing = 1.08
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.space_before = Pt(0)
            if numeric and col_index > 0:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            for part in re.split(r"(\[COMPLETE[^\]]*\])", value):
                r = p.add_run(part)
                r.font.name = "Arial"
                r.font.size = Pt(9.5)
                r.bold = row_index == 0 or part.startswith("[COMPLETE")
                if part.startswith("[COMPLETE"):
                    r.font.color.rgb = RGBColor.from_string("8A4B08")
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "E8EDF2" if row_index == 0 else ("F6F7F8" if row_index % 2 == 0 else "FFFFFF"))
            cell._tc.get_or_add_tcPr().append(shade)
    paragraph("", size=2, after=3)
    return t


footer = section.footer.paragraphs[0]
footer.paragraph_format.space_after = Pt(0)
footer.paragraph_format.tab_stops.add_tab_stop(Inches(6.75), WD_ALIGN_PARAGRAPH.RIGHT)
run = footer.add_run("CONFIDENTIAL | Discussion draft | 6 October 2026\t")
run.font.name = "Arial"
run.font.size = Pt(8)
run.font.color.rgb = RGBColor.from_string("555555")
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), "PAGE")
footer._p.append(field)
footer.add_run(" / 3").font.size = Pt(8)

# PAGE 1
doc.add_paragraph("Circularo Strategic\nShareholder Opportunity", "Title")
paragraph("Prepared for Joshua Jahani by Josef Neumann\nConfidential discussion draft | 6 October 2026", size=9.5, after=12)
paragraph("We are exploring Circularo's first external shareholder transaction to accelerate GCC and international expansion while providing meaningful liquidity to the founders. We would value your perspective on the investor universe, transaction structure and evidence needed to support a market-tested valuation.")

heading("The business today")
paragraph("Circularo is a Dubai-headquartered digital trust software company founded in 2013. Our platform connects document workflows, identity, approvals, digital signatures and transaction evidence for governments and enterprises. The same technology operates as managed SaaS, in customer-controlled infrastructure and as shared services supporting multiple organizations.")
paragraph("We generate recurring revenue through enterprise subscriptions, dedicated platform licensing and partner distribution. We have built the business without external institutional equity funding. Management reports profitable operations and no debt.")

heading("Reported financial performance")
table(
    ["Metric", "2024", "2025"],
    [
        ["Net revenue", "USD 2.41M", "USD 2.79M"],
        ["Adjusted EBITDA", "USD 535K", "USD 859K"],
        ["Adjusted EBITDA margin", "22.2%", "30.7%"],
        ["Gross margin", "66.6%", "69.2%"],
    ],
    [3.15, 1.8, 1.8], numeric=True,
)
paragraph("Figures are management-reported historical results. The adjustment basis and supporting financial statements will be provided for review.", size=9)
paragraph("Our September 2026 subscription snapshot reports 143 active paid commercial accounts across 15 countries. These accounts are distinct from the underlying organizations reached through government platforms and enterprise groups. Revenue remains predominantly GCC-based.")

heading("Current figures to complete")
paragraph("[COMPLETE 1: Latest ARR in USD, snapshot date and definition.]\n[COMPLETE 2: 2026 YTD net revenue and adjusted EBITDA, reporting period, and top-five customer share of ARR.]", size=9.5)

# PAGE 2
doc.add_page_break()
doc.add_paragraph("Growth Through Platforms\nand Distribution", "Heading 1")
paragraph("Our growth strategy builds on positions already established. A government shared-service platform or distribution partner can bring Circularo to multiple organizations through one technology relationship. The investment case depends on converting that reach into paying adoption, increased usage and recurring revenue.")

heading("Existing platforms and expansion opportunities")
paragraph("The following positions are described in management's current briefing. Expansion remains dependent on adoption, procurement and commercial terms; it is not contracted revenue.", size=9)
table(
    ["Platform", "Present position", "Next growth opportunity"],
    [
        ["Digital Dubai\nDigitalSign", "Shared-service deployment with adoption across multiple government entities.", "Broader entity adoption, workflows and potential central procurement."],
        ["TDRA\nGovSign", "Federal shared-service deployment; GovSign 2.0 launched in October 2025.", "Further federal-entity adoption and expanded document processes."],
        ["Sharjah Government\nSharjah Sign", "Shared-service deployment across the Sharjah government ecosystem.", "Additional departments, users and transaction volumes."],
    ],
    [1.45, 2.6, 2.7],
)
paragraph("[COMPLETE 3: For each platform, add current annual Circularo revenue, active versus addressable entities, and the next procurement or adoption milestone with expected timing.]", size=9.5)

heading("Three routes to recurring revenue growth")
paragraph("Direct enterprise adoption grows through new customers and expansion within existing accounts. Government shared services grow through wider paid adoption and additional use cases inside existing platforms, followed by replication in other jurisdictions. Strategic distribution grows through partners commercializing Circularo within their own customer ecosystems.")
paragraph("Management identifies e& DigiSign and TCC/Mokham in Saudi Arabia as examples of the partner model. The economics depend on each agreement: downstream service revenue must be distinguished from the revenue Circularo receives.")
paragraph("[COMPLETE 4: Confirm e& DigiSign and TCC/Mokham commercial status, current Circularo revenue and contracted participation in downstream growth.]", size=9.5)

heading("Strategic value and execution risks")
paragraph("The combination of deployment flexibility, regional identity integrations, enterprise workflows and operating experience in government environments can make Circularo relevant to a buyer seeking GCC market access or complementary trust technology. The strategic value should be tested through specific commercial opportunities the buyer can unlock.")
paragraph("The main risks are GCC and customer concentration, long procurement cycles, and the pace at which platform access converts to paid adoption. International replication and partner distribution also require execution resources. We would welcome a candid assessment of these constraints.", size=10)

# PAGE 3
doc.add_page_break()
doc.add_paragraph("Shareholder Objectives\nand Transaction Discussion", "Heading 1")
paragraph("We want a shareholder who can help turn Circularo's existing technology and customer positions into a materially larger business. Relevant contributions include enterprise distribution, new shared-service platforms, geographic expansion and integration into a broader technology offering.")

heading("The transaction we want to explore")
table(
    ["Objective", "Starting position for discussion"],
    [
        ["Founder liquidity", "Meaningful secondary proceeds are required."],
        ["Ownership", "A minority interest of up to approximately 25% is preferred."],
        ["Indicative size", "USD 10-15M is a working scenario to test, not a fixed offer."],
        ["Primary capital", "Optional where a defined growth plan justifies it."],
        ["Valuation and structure", "To be tested against buyer appetite, evidence and strategic contribution."],
    ],
    [1.6, 5.15],
)
paragraph("A wholly secondary purchase of 25% for USD 10-15M would imply USD 40-60M equity value. We would value your assessment of whether the business and buyer-specific opportunities support that expectation, and which alternative structures merit consideration.", size=9.5)
paragraph("[COMPLETE 5: Minimum desired founder cash proceeds and any primary capital requirement for the agreed growth plan.]", size=9.5)

heading("The shareholder we are looking for")
paragraph("Our preferred routes include a strategic technology company, a sovereign or institutional investor, or a group with relevant GCC and international reach. We are open to considering financial investors with demonstrable commercial capabilities. We would like to understand which candidates can provide founder liquidity and make a measurable contribution to growth.")
paragraph("Any strategic rationale should translate into concrete initiatives, such as distribution resources, portfolio deployments, integrations or named market-entry programmes. We also want to understand the governance and alignment required to make those initiatives work.")

heading("Longer term product direction")
paragraph("We are developing richer trusted records and deeper authority controls. Our longer-term vision is to govern and evidence consequential actions initiated by enterprise systems and AI agents. This is development and strategic optionality; it is separate from current revenue and released product capability.", size=10)

heading("Questions for our discussion")
for number, text in enumerate([
    "Which buyer categories are credible at our scale and willing to support meaningful founder secondary?",
    "What evidence and comparable transactions should determine valuation, and how does the minority preference affect appetite?",
    "Which commercial initiatives could justify strategic ownership, and which could be achieved through a partnership?",
    "What preparation, execution team and engagement economics would you recommend for a focused process?",
], 1):
    paragraph(f"{number}. {text}", size=10, after=5)

doc.save(OUTPUT)
print(f"Created {OUTPUT}")
