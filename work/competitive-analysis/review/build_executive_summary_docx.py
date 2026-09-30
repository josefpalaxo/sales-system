"""Editable Word counterpart to the three-page executive summary PDF."""
from pathlib import Path
import re, struct, uuid
from zipfile import ZipFile, ZIP_DEFLATED
from lxml import etree
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT=Path(__file__).resolve().parents[3]
OUT=ROOT/'work/competitive-analysis'
ASSETS=ROOT/'.agents/skills/circularo-slides/assets'
DEST=OUT/'Circularo-Executive-Competitive-Summary.docx'
source_text=(OUT/'Circularo-Executive-Competitive-Summary.md').read_text()
source_rows=re.findall(r'^(\d+)\. \[([^\]]+)\]\((https?://[^)]+)\)\.',source_text,re.M)
source_map={n:url for n,_,url in source_rows}
NAVY='1D0090';PURPLE='7000FF';BODY='3D3D3D';MUTED='667085';LINE='E4E7EC';LIGHT='F5EDFF'
WIDTH=515.276
doc=Document();sec=doc.sections[0]
sec.page_width=Pt(595.276);sec.page_height=Pt(841.89)
sec.left_margin=Pt(46);sec.right_margin=Pt(34);sec.top_margin=Pt(83);sec.bottom_margin=Pt(44)
sec.header_distance=Pt(31);sec.footer_distance=Pt(17)
doc.core_properties.title='Circularo | Executive Competitive Summary'
doc.core_properties.subject='Unified Sovereign Trust Platform vs. Alternatives'
doc.core_properties.author='Circularo';doc.core_properties.last_modified_by='Circularo'

def set_font(target,name='Mulish',size=10.2,bold=False,color=BODY):
    target.font.name=name;target.font.size=Pt(size);target.font.bold=bold;target.font.color.rgb=RGBColor.from_string(color)
    rp=target._element.get_or_add_rPr()
    rf=rp.find(qn('w:rFonts'))
    if rf is None:rf=OxmlElement('w:rFonts');rp.insert(0,rf)
    for k in list(rf.attrib):
        if k.lower().endswith('theme'):del rf.attrib[k]
    for k in ('ascii','hAnsi','eastAsia','cs'):rf.set(qn('w:'+k),name)

def style(name,font,size,leading,color=BODY,bold=False,before=0,after=0):
    s=doc.styles[name] if name in doc.styles else doc.styles.add_style(name,1)
    set_font(s,font,size,bold,color)
    f=s.paragraph_format;f.space_before=Pt(before);f.space_after=Pt(after);f.line_spacing=Pt(leading)
    f.right_indent=Pt(12);f.widow_control=False
    return s
style('Normal','Mulish',10.2,14.1,after=8)
style('Title','Spartan',25.5,32,NAVY,True,after=13)
style('Heading 1','Spartan',22,28,NAVY,True,after=10)
style('Heading 2','Spartan',12.2,16,NAVY,True,before=12,after=6)
style('Kicker','Mulish',9,13,PURPLE,True,after=11)
style('Cell','Mulish',9.15,12.3)
style('CellHeader','Mulish',9.2,12,'FFFFFF',True)
style('Note','Mulish',8.1,10.8,MUTED,before=7,after=7)
style('Sources','Mulish',7.7,10.2,MUTED,after=2)
style('SourceHeading','Mulish',8,11,NAVY,True,before=8,after=5)
for name in ['Title','Heading 1','Heading 2','SourceHeading']:doc.styles[name].paragraph_format.keep_with_next=True
for s in doc.styles:
    for b in list(s.element.findall('.//'+qn('w:pBdr'))):b.getparent().remove(b)
for name in ['Header','Footer']:
    doc.styles[name].paragraph_format.tab_stops.clear_all()

def line_border(p,side,color=LINE,space=0,size=4):
    pr=p._p.get_or_add_pPr();b=pr.find(qn('w:pBdr'))
    if b is None:b=OxmlElement('w:pBdr');pr.append(b)
    e=OxmlElement('w:'+side)
    for k,v in [('val','single'),('sz',str(size)),('space',str(space)),('color',color)]:e.set(qn('w:'+k),v)
    b.append(e)

def link(p,text,url,size,color=NAVY,bold=False,super_=False):
    rid=p.part.relate_to(url,RT.HYPERLINK,is_external=True)
    h=OxmlElement('w:hyperlink');h.set(qn('r:id'),rid)
    r=OxmlElement('w:r');rp=OxmlElement('w:rPr')
    fonts=OxmlElement('w:rFonts')
    for k in ['ascii','hAnsi','cs']:fonts.set(qn('w:'+k),'Mulish')
    rp.append(fonts)
    for tag,val in [('color',color),('sz',str(round(size*2))),('szCs',str(round(size*2)))]:
        e=OxmlElement('w:'+tag);e.set(qn('w:val'),val);rp.append(e)
    if bold:rp.append(OxmlElement('w:b'))
    if super_:
        v=OxmlElement('w:vertAlign');v.set(qn('w:val'),'superscript');rp.append(v)
    r.append(rp);t=OxmlElement('w:t');t.text=text;t.set(qn('xml:space'),'preserve');r.append(t);h.append(r);p._p.append(h)

def add_text(p,text,style_name=None):
    if style_name:p.style=doc.styles[style_name]
    if text.startswith('*') and not text.startswith('**') and text.endswith('*'):text=text[1:-1]
    bold=False
    for token in re.split(r'(\*\*|(?:\[\d+\])+)',text):
        if token=='**':bold=not bold;continue
        if re.fullmatch(r'(?:\[\d+\])+',token or ''):
            ns=re.findall(r'\d+',token)
            for i,n in enumerate(ns):
                if i:
                    r=p.add_run(', ');set_font(r,size=10,color=PURPLE);r.font.superscript=True
                link(p,n,source_map[n],10,PURPLE,super_=True)
        elif token:
            r=p.add_run(token)
            if bold:r.bold=True
    return p

def paragraph(text,style_name='Normal'):
    return add_text(doc.add_paragraph(style=style_name),text)

def spacer(size):
    p=doc.add_paragraph();f=p.paragraph_format;f.space_before=f.space_after=Pt(0);f.line_spacing=Pt(size)
    p.add_run().font.size=Pt(1)

def cell_props(cell,fill=None,padding=9,bottom_border=True):
    pr=cell._tc.get_or_add_tcPr();m=OxmlElement('w:tcMar')
    for side in ['top','bottom','left','right']:
        e=OxmlElement('w:'+side);e.set(qn('w:w'),str(int(padding*20)));e.set(qn('w:type'),'dxa');m.append(e)
    pr.append(m)
    if fill:
        sh=OxmlElement('w:shd');sh.set(qn('w:fill'),fill);pr.append(sh)
    borders=OxmlElement('w:tcBorders')
    for side in ['top','left','bottom','right']:
        e=OxmlElement('w:'+side);e.set(qn('w:val'),'single' if side=='bottom' and bottom_border else 'nil');e.set(qn('w:sz'),'3');e.set(qn('w:color'),LINE);borders.append(e)
    pr.append(borders);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP

def base_table(container,rows,widths,**kw):
    t=container.add_table(rows=rows,cols=len(widths),**kw);t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.LEFT
    for col,w in zip(t.columns,widths):col.width=Pt(w)
    for row in t.rows:
        for cell,w in zip(row.cells,widths):cell.width=Pt(w)
        trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
    indent=OxmlElement('w:tblInd');indent.set(qn('w:w'),'0');indent.set(qn('w:type'),'dxa');t._tbl.tblPr.append(indent)
    return t

def content_table(lines):
    rows=[[v.strip() for v in l.strip().strip('|').split('|')] for l in lines]
    rows=[r for r in rows if not all(re.fullmatch(r'[:\- ]+',v) for v in r)]
    compare=len(rows[0])==4;widths=[88,157,135,WIDTH-380] if compare else [108,WIDTH-108]
    t=base_table(doc,len(rows),widths)
    t._tbl.tblPr.find(qn('w:tblInd')).set(qn('w:w'),'180')
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for i,row in enumerate(t.rows):
        for j,cell in enumerate(row.cells):
            fill=NAVY if i==0 else LIGHT if compare and j==1 else 'F8F9FC' if i%2 else 'FFFFFF'
            cell_props(cell,fill)
            p=cell.paragraphs[0];add_text(p,rows[i][j],'CellHeader' if i==0 else 'Cell')
            p.paragraph_format.right_indent=Pt(0);p.paragraph_format.keep_with_next=False
    spacer(8)

def model():
    gap=17;w=(WIDTH-gap*3)/4
    t=base_table(doc,1,[w,gap,w,gap,w,gap,w])
    t._tbl.tblPr.find(qn('w:tblInd')).set(qn('w:w'),'60')
    labels=[('ONE OPERATOR','Central service operation'),('SHARED PLATFORM','Common trust infrastructure'),('PARTICIPATING ENTITIES','Organization-specific controls'),('PROCESSES & EVIDENCE','Approval, signing and records')]
    for i,cell in enumerate(t.rows[0].cells):
        cell_props(cell,LIGHT if i%2==0 else 'FFFFFF',3,False)
        cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p=cell.paragraphs[0];p.paragraph_format.right_indent=Pt(0);p.paragraph_format.line_spacing=Pt(10);p.paragraph_format.space_after=Pt(3)
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        if i%2:
            r=p.add_run('→');set_font(r,size=12,color=PURPLE)
        else:
            title,caption=labels[i//2];r=p.add_run(title);set_font(r,size=7.8,bold=True,color=NAVY)
            p2=cell.add_paragraph();p2.alignment=WD_ALIGN_PARAGRAPH.CENTER
            f=p2.paragraph_format;f.right_indent=Pt(0);f.line_spacing=Pt(10);f.space_after=f.space_before=Pt(0)
            r=p2.add_run(caption);set_font(r,size=7.8)
    spacer(8)

# Same recurring header and footer as the PDF.
sec.header_distance=Pt(27)
header_blank=sec.header.paragraphs[0]
ht=base_table(sec.header,1,[120,WIDTH-120],width=Pt(WIDTH))
ht._tbl.tblPr.find(qn('w:tblInd')).set(qn('w:w'),'-120')
for cell in ht.rows[0].cells:
    cell_props(cell,None,0,True);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    cell._tc.get_or_add_tcPr().find(qn('w:tcMar')).find(qn('w:bottom')).set(qn('w:w'),'140')
    p=cell.paragraphs[0];p.style=doc.styles['Normal'];p.paragraph_format.right_indent=Pt(0)
    p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.0
p=ht.cell(0,0).paragraphs[0]
p.add_run().add_picture(str(ASSETS/'logos/circularo-logo-logotype-blue.png'),width=Pt(99))
p=ht.cell(0,1).paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=p.add_run('PARTNER & INVESTOR PERSPECTIVE');set_font(r,size=8,bold=True,color=MUTED)
header_blank._p.getparent().remove(header_blank._p)
footer=sec.footer.paragraphs[0];f=footer.paragraph_format
f.left_indent=Pt(-6);f.right_indent=Pt(6);f.space_after=Pt(0);f.line_spacing=Pt(10)
f.tab_stops.add_tab_stop(Pt(WIDTH-6),WD_TAB_ALIGNMENT.RIGHT);line_border(footer,'top',space=6)
r=footer.add_run('CIRCULARO  /  SEPTEMBER 2026\t');set_font(r,size=7.6,color=MUTED)
for field in ['PAGE','NUMPAGES']:
    r=footer.add_run();set_font(r,size=7.6,color=MUTED)
    begin=OxmlElement('w:fldChar');begin.set(qn('w:fldCharType'),'begin');r._r.append(begin)
    instr=OxmlElement('w:instrText');instr.text=field;r._r.append(instr)
    end=OxmlElement('w:fldChar');end.set(qn('w:fldCharType'),'end');r._r.append(end)
    if field=='PAGE':r=footer.add_run(' / ');set_font(r,size=7.6,color=MUTED)

lines=source_text.split('### Selected sources')[0].splitlines();i=0
while i<len(lines):
    line=lines[i].strip()
    if not line or line=='# Circularo':i+=1;continue
    if line.startswith('## Unified'):
        paragraph('Unified Sovereign Trust\nPlatform vs. Alternatives','Title');i+=1;continue
    if line.startswith('**Executive Competitive Summary'):
        paragraph('EXECUTIVE COMPETITIVE SUMMARY','Kicker');i+=1;continue
    if line.startswith('### '):
        title=line[4:]
        if title in ('How Circularo compares','Why this matters to partners'):
            p=paragraph('Partner & investment perspective' if title=='Why this matters to partners' else title,'Heading 1')
            p.paragraph_format.page_break_before=True
            if title=='Why this matters to partners':paragraph(title,'Heading 2')
        else:paragraph(title,'Heading 2')
        i+=1;continue
    if line.startswith('|'):
        group=[]
        while i<len(lines) and lines[i].strip().startswith('|'):group.append(lines[i]);i+=1
        content_table(group);continue
    group=[line];i+=1
    while i<len(lines) and lines[i].strip() and not lines[i].strip().startswith(('#','|')):
        group.append(lines[i].strip());i+=1
    text=' '.join(group)
    if text.startswith('**One operator'):model()
    else:paragraph(text,'Note' if text.startswith('*Comparison basis:') else 'Normal')

paragraph('SELECTED SOURCES  /  CLICK TO OPEN','SourceHeading')
t=base_table(doc,1,[WIDTH/2,WIDTH/2]);mid=(len(source_rows)+1)//2
for cell,rows in zip(t.rows[0].cells,[source_rows[:mid],source_rows[mid:]]):
    cell_props(cell,None,0,False)
    for j,(n,title,url) in enumerate(rows):
        p=cell.paragraphs[0] if j==0 else cell.add_paragraph();p.style=doc.styles['Sources']
        p.paragraph_format.right_indent=Pt(12)
        r=p.add_run(n+'. ');r.bold=True
        short=title.replace('Circularo: ','Circularo / ').replace('Adobe Acrobat Sign: ','Adobe / ').replace('Namirial eSignAnyWhere: ','Namirial / ')
        link(p,short,url,7.7,NAVY)

# Embed the actual brand fonts so the editable layout travels with the document.
for setting in ['embedTrueTypeFonts','embedSystemFonts']:
    e=OxmlElement('w:'+setting);doc.settings.element.append(e)
doc.save(DEST)
NS={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
REL='http://schemas.openxmlformats.org/package/2006/relationships'
with ZipFile(DEST) as z:parts={n:z.read(n) for n in z.namelist()}
font_table=etree.fromstring(parts['word/fontTable.xml'])
rels=etree.Element('{'+REL+'}Relationships',nsmap={None:REL})
for idx,(family,kind,filename) in enumerate([('Mulish','embedRegular','mulish-regular.ttf'),('Mulish','embedBold','mulish-bold.ttf'),('Spartan','embedBold','spartan-bold.ttf')],1):
    raw=bytearray((ASSETS/'fonts'/filename).read_bytes())
    for j in range(struct.unpack('>H',raw[4:6])[0]):
        base=12+16*j
        if raw[base:base+4]==b'OS/2':
            offset=struct.unpack('>I',raw[base+8:base+12])[0]
            assert not struct.unpack('>H',raw[offset+8:offset+10])[0]&2,'Restricted font embedding'
    key=uuid.uuid5(uuid.NAMESPACE_URL,'circularo-executive/'+filename)
    mask=key.bytes[::-1]
    for j in range(32):raw[j]^=mask[j%16]
    path=f'fonts/brand-{idx}.odttf';rid=f'rIdBrandFont{idx}'
    parts['word/'+path]=bytes(raw)
    font=font_table.find("w:font[@w:name='"+family+"']",NS)
    if font is None:font=etree.SubElement(font_table,qn('w:font'));font.set(qn('w:name'),family)
    e=etree.SubElement(font,qn('w:'+kind));e.set(qn('r:id'),rid);e.set(qn('w:fontKey'),'{'+str(key).upper()+'}');e.set(qn('w:subsetted'),'false')
    e=etree.SubElement(rels,'{'+REL+'}Relationship');e.set('Id',rid);e.set('Type',NS['r']+'/font');e.set('Target',path)
parts['word/fontTable.xml']=etree.tostring(font_table,xml_declaration=True,encoding='UTF-8',standalone=True)
parts['word/_rels/fontTable.xml.rels']=etree.tostring(rels,xml_declaration=True,encoding='UTF-8',standalone=True)
ct=etree.fromstring(parts['[Content_Types].xml'])
e=etree.SubElement(ct,'{http://schemas.openxmlformats.org/package/2006/content-types}Default');e.set('Extension','odttf');e.set('ContentType','application/vnd.openxmlformats-officedocument.obfuscatedFont')
parts['[Content_Types].xml']=etree.tostring(ct,xml_declaration=True,encoding='UTF-8',standalone=True)
with ZipFile(DEST,'w',ZIP_DEFLATED) as z:
    for n,data in parts.items():z.writestr(n,data)
print(DEST)
