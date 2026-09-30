import hashlib,json,subprocess,importlib.util
from pathlib import Path
from decimal import Decimal as D
from openpyxl import load_workbook
from zipfile import ZipFile
from lxml import etree
root=Path(__file__).resolve().parents[1];out=root/'outputs/dda-business-case'
original=root/'DDA deparments shared service.xlsx'
source=json.loads((root/'build/source-data.json').read_text())
assert hashlib.sha256(original.read_bytes()).hexdigest()==source['source_sha256']
model=out/'DDA-business-case-model.xlsx';w=load_workbook(model,data_only=True);wf=load_workbook(model,data_only=False)
errors=[(s.title,c.coordinate,c.value) for s in w for row in s for c in row if c.data_type=='e'];assert not errors,errors
onprem=D(1260);baseline=[];central=[]
for u,n in zip([10000,10625,11250,11875,12500],[10,6,8,8,8]):
    b=D(u)*D(2360)+(D(u)-onprem)*D(236)+onprem*D(472)
    c=D(u)*D(1180)+(D(u)-onprem)*D(118)+onprem*D(236)+D(n)*D(18000)
    baseline.append(b);central.append(c)
assert w['Cost model']['K23'].value==float(sum(baseline))
assert w['Cost model']['K32'].value==float(sum(central))
assert w['Cost model']['K34'].value==float(sum(baseline)-sum(central))
assert len(wf['Executive']._charts)==2
assert w['Assumptions']['E4'].value==1
assert len(source['entities'])==92
for i,e in enumerate(source['entities'],8):assert w['Entity scope'].cell(i,5).value==e['name']
for i,s in enumerate(source['survey'],35):
    assert w['Demand'].cell(i,5).value==s['count']
    assert w['Demand'].cell(i,6).value==s['percent']
assert all(w['Customers'].cell(i,9).value is None for i in [10,11,12,13,14,16])
assert w['Customers']['H15'].value is None
doc=ZipFile(out/'DDA-executive-brief.docx');assert b'pBdr' not in doc.read('word/document.xml')
presentation=ZipFile(out/'DDA-executive-slides.pptx')
slides=[n for n in presentation.namelist() if n.startswith('ppt/slides/slide') and n.endswith('.xml')]
assert len(slides)==3
ns={'p':'http://schemas.openxmlformats.org/presentationml/2006/main','a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
for p in slides:
    xml=etree.fromstring(presentation.read(p));pics=xml.xpath('//p:pic',namespaces=ns)
    assert len(pics)==1
    off=pics[0].xpath('.//a:xfrm/a:off',namespaces=ns)[0]
    ext=pics[0].xpath('.//a:xfrm/a:ext',namespaces=ns)[0]
    assert (int(off.get('x')),int(off.get('y')),int(ext.get('cx')),int(ext.get('cy')))==(457200,228600,381000,381000)
report={'source_unchanged':True,'formula_errors':errors,'base_baseline':float(sum(baseline)),'base_cost':float(sum(central)),'base_net':float(sum(baseline)-sum(central)),'source_rows':92,'survey_rows':47,'native_excel_charts':2,'slides':3,'logo_geometry':'40 x 40 px at 48,24 on all slides','native_office_checked':False}
(root/'build/verification.json').write_text(json.dumps(report,indent=2));print(json.dumps(report,indent=2))
