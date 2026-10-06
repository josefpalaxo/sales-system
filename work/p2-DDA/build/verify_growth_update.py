import os, json, math, subprocess, hashlib, shutil, re
from copy import copy
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from xml.etree import ElementTree as ET
from openpyxl import load_workbook

root=Path(__file__).resolve().parents[1]
qa=root/'build/growth-50'
final=root/'outputs/dda-business-case-v2/DDA-savings-model-v2.2xlsx.xlsx'
candidate=qa/final.name
before=load_workbook(qa/'before.xlsx',data_only=False)
after=load_workbook(candidate,data_only=False)
assert before.sheetnames==after.sheetnames
value_changes=[];style_changes=[]
for sheet in before:
    target=after[sheet.title]
    assert list(sheet.merged_cells.ranges)==list(target.merged_cells.ranges)
    assert len(sheet._charts)==len(target._charts)
    assert str(sheet.data_validations)==str(target.data_validations)
    for row in sheet:
        for c in row:
            d=target[c.coordinate]
            if c.value!=d.value:value_changes.append((sheet.title,c.coordinate,c.value,d.value))
            if any(copy(getattr(c,k))!=copy(getattr(d,k)) for k in ['font','fill','border','alignment','number_format','protection']):
                style_changes.append((sheet.title,c.coordinate))
assert value_changes==[('Setup','E67',.2,.5)],value_changes
print(json.dumps({'value_changes':value_changes,'style_changes':len(style_changes),'style_examples':style_changes[:8]}))

native=qa/'native';native.mkdir(exist_ok=True)
profile=qa/'profile';profile.mkdir(exist_ok=True)
tmp=qa/'tmp';tmp.mkdir(exist_ok=True)
env=os.environ.copy();env['TMPDIR']=str(tmp)
for key,folder in [('XDG_CONFIG_HOME','config'),('XDG_CACHE_HOME','cache')]:
    d=profile/folder;d.mkdir(exist_ok=True);env[key]=str(d)
soffice='/Users/josefneumann/.cache/codex-runtimes/codex-primary-runtime/dependencies/bin/override/soffice'
run=subprocess.run([soffice,f'-env:UserInstallation={profile.as_uri()}','--headless','--norestore','--convert-to','xlsx','--outdir',str(native),str(candidate)],env=env,capture_output=True,text=True,timeout=55)
assert run.returncode==0,run.stderr
book=load_workbook(native/final.name,data_only=True)
errors=[(s.title,c.coordinate,c.value)for s in book for row in s for c in row if c.data_type=='e']
assert not errors,errors
existing=1965;survey=4921;other=3114;forecast=[]
for i,col in enumerate('DEFGH'):
    if i:existing,survey,other=[math.floor(x*1.05+.5)for x in (existing,survey,other)]
    paid=max(10000,existing+survey+other)
    benchmark=1929169.75*existing/1310+(survey+other)*2950
    subscription=paid*1180
    values={28:existing,29:survey,30:other,32:paid,38:benchmark,39:subscription,40:benchmark-subscription}
    for row,value in values.items():assert abs(book['Potential'][f'{col}{row}'].value-value)<.01,(col,row)
    forecast.append(values)
for row in (38,39,40):assert abs(book['Potential'][f'I{row}'].value-sum(x[row]for x in forecast))<.01
assert book['Setup']['E67'].value==.5
assert abs(book['Summary']['E32'].value-.6886)<1e-10
assert abs(book['Potential']['E100'].value)<.01

# Artifact Tool preserved every original cell value/formula except E67, all
# effective styles, merged ranges, validations and native chart counts.
assert not style_changes,style_changes[:10]
for sheet in before:
    target=after[sheet.title]
    assert sheet.freeze_panes==target.freeze_panes
    assert str(sheet.column_dimensions)==str(target.column_dimensions)
    assert str(sheet.row_dimensions)==str(target.row_dimensions)
    for a,b in zip(sheet._charts,target._charts):
        assert ET.tostring(a.to_tree())==ET.tostring(b.to_tree())
artifact_values=load_workbook(candidate,data_only=True)
for sheet in after:
    for row in sheet:
        for c in row:
            if c.data_type!='f':continue
            x=artifact_values[sheet.title][c.coordinate].value;y=book[sheet.title][c.coordinate].value
            if isinstance(x,(int,float)) and isinstance(y,(int,float)):assert abs(x-y)<.01,(sheet.title,c.coordinate,x,y)
            else:assert x==y,(sheet.title,c.coordinate,x,y)
shutil.copy2(candidate,final)
saved=load_workbook(final,data_only=True)
assert saved['Potential']['E23'].value==655
assert abs(saved['Summary']['H13'].value-book['Potential']['I40'].value)<.01
prior=json.loads((root/'build/dda-executive/model-data.json').read_text())
prior['sha256']=hashlib.sha256(final.read_bytes()).hexdigest()
for name,cells in prior['cells'].items():
    for address in cells:cells[address]=saved[name][address].value
prior['customers']=[[saved['Existing customers'].cell(r,c).value for c in range(3,13)]for r in range(9,16)]
(root/'build/dda-executive/model-data.json').write_text(json.dumps(prior,indent=2))
(qa/'verification.json').write_text(json.dumps({'changed_input':'Setup!E67','new_growth':.5,'all_original_cell_styles_preserved':True,'native_formula_errors':errors,'forecast':forecast,'five_year_savings':saved['Potential']['I40'].value},indent=2))
print(json.dumps({'published':str(final),'five_year_savings':saved['Potential']['I40'].value}))
