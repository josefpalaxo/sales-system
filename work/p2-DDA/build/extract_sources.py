import json, re, hashlib
from pathlib import Path
from openpyxl import load_workbook

root = Path(__file__).resolve().parents[1]
book = root / 'DDA deparments shared service.xlsx'
wb = load_workbook(book, read_only=True, data_only=True)
entities = []
notes = {}
for line in (root/'entity-research-register.md').read_text().splitlines():
    cells = [x.strip() for x in line.split('|')[1:-1]]
    if len(cells)==6 and cells[0].isdigit(): notes[int(cells[0])] = cells
for row, values in enumerate(wb['dda departments'].iter_rows(min_row=2,max_row=93,min_col=1,max_col=2,values_only=True),2):
    n = notes[row]
    entities.append({'row':row,'number':values[0],'name':values[1],'workforce':n[3],'existing':None if n[4]=='—' else int(n[4]),'note':n[5]})
survey=[]
question=''
for line in (root/'fact-register.md').read_text().split('## DDA internal entity survey')[1].split('### Survey interpretation')[0].splitlines():
    if line.startswith('### Q'): question=line[4:]
    c=[x.strip() for x in line.split('|')[1:-1]]
    if len(c)==3 and c[1].isdigit(): survey.append({'question':question,'response':c[0],'count':int(c[1]),'percent':None if c[2]=='—' else float(c[2].strip('%'))/100})
sources=[]
for line in (root/'research-notes.md').read_text().splitlines():
    if re.match(r'\| R\d\d \|',line):
        c=[x.strip() for x in line.split('|')[1:-1]]
        urls=re.findall(r'\]\((https?://[^)]+)\)',c[1])
        sources.append({'id':c[0],'citation':c[1],'finding':c[2],'limits':c[3],'url':urls[0] if urls else ''})
(root/'build/source-data.json').write_text(json.dumps({'entities':entities,'survey':survey,'research':sources,'source_sha256':hashlib.sha256(book.read_bytes()).hexdigest()},indent=2))
print(f'Extracted {len(entities)} source rows, {len(survey)} survey rows and {len(sources)} research sources; source unchanged.')
