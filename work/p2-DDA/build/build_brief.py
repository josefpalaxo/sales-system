import json
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT

root=Path(__file__).resolve().parents[1]
m=json.loads((root/'build/model-results.json').read_text())
doc=Document(); sec=doc.sections[0]
sec.page_width=Inches(8.27); sec.page_height=Inches(11.69)
sec.top_margin=sec.bottom_margin=Inches(.67);sec.left_margin=sec.right_margin=Inches(.72)
sec.header_distance=sec.footer_distance=Inches(.28)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Header','Footer']:
    st=doc.styles[name];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0)
doc.styles['Normal'].font.size=Pt(10.5)
doc.styles['Normal'].paragraph_format.space_after=Pt(7)
doc.styles['Normal'].paragraph_format.line_spacing=1.12
doc.styles['Title'].font.size=Pt(25);doc.styles['Title'].font.bold=True
doc.styles['Title'].paragraph_format.space_after=Pt(11)
doc.styles['Subtitle'].font.size=Pt(11)
for n,size in [('Heading 1',16),('Heading 2',12)]:
    doc.styles[n].font.size=Pt(size);doc.styles[n].font.bold=True
    doc.styles[n].paragraph_format.space_before=Pt(12);doc.styles[n].paragraph_format.space_after=Pt(7)
hp=sec.header.paragraphs[0];hp.text='CIRCULARO   /   DIGITAL DUBAI AUTHORITY';hp.runs[0].font.size=Pt(8)
ft=sec.footer.paragraphs[0];ft.alignment=WD_ALIGN_PARAGRAPH.RIGHT
ft.add_run('DDA central procurement   |   Internal review draft   |   ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');ft._p.append(field)
for r in ft.runs:r.font.size=Pt(8)
def p(t,bold=False):
    q=doc.add_paragraph();q.add_run(t).bold=bold;return q
def h(t):doc.add_heading(t,1)
def money(v):return f'{v/1e6:,.2f}'
def table(headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for c,w in zip(t.columns,widths):c.width=Inches(w)
    for c,txt in zip(t.rows[0].cells,headers):c.text=txt
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for v in rows:
        for c,x in zip(t.add_row().cells,v):c.text=str(x)
    for ri,r in enumerate(t.rows):
        for ci,c in enumerate(r.cells):
            c.width=Inches(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=c._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'1D0090' if ri==0 else ('F3F5F8' if ri%2 else 'FFFFFF'));pr.append(sh)
            margins=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                v=OxmlElement('w:'+side);v.set(qn('w:w'),'95');v.set(qn('w:type'),'dxa');margins.append(v)
            pr.append(margins)
            borders=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                v=OxmlElement('w:'+side);v.set(qn('w:val'),'single');v.set(qn('w:sz'),'4');v.set(qn('w:color'),'D9D9D9');borders.append(v)
            pr.append(borders)
            for q in c.paragraphs:
                q.paragraph_format.space_after=Pt(0);q.paragraph_format.line_spacing=1.05
                q.alignment=WD_ALIGN_PARAGRAPH.LEFT if ci==0 or len(headers)==2 else WD_ALIGN_PARAGRAPH.RIGHT
                for run in q.runs:
                    run.font.size=Pt(10);run.bold=ri==0
                    if ri==0:run.font.color.rgb=RGBColor(255,255,255)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)
    return t
def link(label,url):
    q=doc.add_paragraph();rel=q.part.relate_to(url,RT.HYPERLINK,is_external=True)
    hyper=OxmlElement('w:hyperlink');hyper.set(qn('r:id'),rel);r=OxmlElement('w:r');pr=OxmlElement('w:rPr');sz=OxmlElement('w:sz');sz.set(qn('w:val'),'18');pr.append(sz);r.append(pr);text=OxmlElement('w:t');text.text=label;r.append(text);hyper.append(r);q._p.append(hyper)

doc.add_paragraph('DDA central procurement business case',style='Title')
doc.add_paragraph('Circularo Digital Sovereign Sign\nExecutive decision brief   20 September 2026',style='Subtitle')
p('Recommendation',True)
p('Proceed with a centrally funded Circularo procurement for an initial 10,000-user annual commitment over five years. The proposed subscription rate of AED 1,180 per user per year halves the agreed DDA Shared Service baseline of AED 2,360. Complete funded entity allocations and the outstanding commercial details before making the commitment unconditional.')
h('Financial case')
p(f'The base planning case produces AED {money(m["totals"]["34"])}m of five-year comparative benefit after modeled onboarding, against AED {money(m["totals"]["23"])}m for equivalent capacity at the DDA baseline. Central modeled cost is AED {money(m["totals"]["32"])}m. This is avoided cost for the expanded service, rather than a claim about current government expenditure.')
table(['Period','Paid users','Baseline\nAED m','Central\nAED m','Benefit\nAED m'],[[f'Year {i+1}',f'{m["annual"]["8"][i]:,}',money(m['annual']['23'][i]),money(m['annual']['32'][i]),money(m['annual']['34'][i])]for i in range(5)]+[['Five years','56,250 user-years',money(m['totals']['23']),money(m['totals']['32']),money(m['totals']['34'])]],[1.04,1.56,1.12,1.12,1.12])
p('The base case assumes 25% cumulative growth, evenly reaching 12,500 users in Year 5, 1,260 existing on-premise users retained, and 40 new entities onboarded over five years. The growth timing, allocation and onboarding schedule are planning assumptions. Added seats use AED 1,180, subject to confirmation of their price protection.')
p(f'Five-year procurement ROI is {m["roi"]:.1%}, defined as net comparative benefit divided by central modeled program cost. Net savings equal 49.5% of the baseline after onboarding. These undiscounted measures exclude unpriced services and tax. They do not include productivity benefits.')
h('Initial budget')
p(f'Year 1 subscription and support total AED 13,128,680 with the assumed deployment mix. Ten new entities add AED 180,000, bringing the modeled Year 1 budget to AED 13,308,680. The initial 10,000-seat recurring commitment totals AED {money(m["initialPoolCost"])}m over five years before growth and onboarding. The all-hosted AED 12.98m annual reference excludes the on-premise support uplift. Regular Basic Price of AED 2,950 provides a 60% reference discount only.')

doc.add_page_break()
h('Demand for the central service')
p('The case combines an existing customer base, reported survey demand and a wider potential entity population. Each provides a different kind of evidence. They cannot be added together until DDA maps respondents and confirms entity allocations.')
table(['Evidence','What it establishes'],[
['Existing customers','1,830 project users across eight entities and nine projects, including RTA’s 1,000 CAPEX users.'],
['DDA internal survey','38 responses on current use. Of 22 user-count answers, 21 numerical bands imply at least 4,921 current or expected users.'],
['New demand','10 of 16 interested organizations report a current requirement, and six report possible future need. Interest is not a purchase commitment.'],
['Expected growth','17 of 22 respondents expect usage to increase. The survey supports growth direction, without specifying 25% or a timetable.'],
['Wider population','92 source-name rows, 91 numbered, with duplicates and aliases. A final count of distinct eligible entities remains unresolved.']],[1.42,4.64])
p('A transparent sensitivity assigns midpoints to bounded user bands and 2,000 users to each of the four “more than 1,000” respondents. It implies 10,558.5 users, excluding the unknown answer. This illustrates why 10,000 merits validation, but does not establish exact demand or incremental users beyond existing customers.')
p('Digital Dubai reported that Smart Employee served over 76,000 employees across 76 government entities in May 2025. The initial pool is about 13.2% of the 76,000 reference denominator. This is evidence of ecosystem scale, not a signature-license conversion rate. Unlimited external recipients also mean external signers should not automatically count as paid users.')
h('Existing customer transition')
p(f'The supplied existing customer records contain AED {m["currentKnown"]:,.2f} in known annual charges. Applying the new subscription and support terms to their 1,830 project users produces AED {m["currentCentral"]:,.2f}. The conditional increase of AED {m["currentCentral"]-m["currentKnown"]:,.2f} shows why the portfolio benefit must rest on expansion at lower benchmark prices. Current amounts assume AED and additive support; missing support and tax details remain unresolved.')
p('RTA is the largest transition difference: its reported AED 608,000 annual support becomes AED 1,416,000 of subscription and 20% support, a conditional increase of AED 808,000. Historic CAPEX is not a recurring saving. Other entities can also differ from the benchmark because their current commercial packages vary.')
p('All four existing SaaS customers qualify for charge cessation upon migration, as does DDA Shared Service. For on-premise projects, 20% support replaces standard 10% support and old support charges. Every deployment receives the lower subscription unit rate. Existing entities incur no new-entity onboarding fee.')

doc.add_page_break()
h('Five year price protection')
p('The central offer fixes subscription pricing for five years, with annual billing and a five-year commitment. This gives DDA unit-cost predictability even when independent prices remain flat. Quantities can grow, so a fixed unit price does not mean a fixed annual budget.')
p(f'The headline case assumes zero baseline price escalation. A separate, explicitly hypothetical 3% annual escalation sensitivity adds AED {money(m["totals"]["46"])}m of avoided baseline increases, taking comparative benefit to AED {money(m["totals"]["47"])}m. This sensitivity is not a guaranteed saving and is excluded from the AED {money(m["totals"]["34"])}m headline.')
h('Implementation and budget control')
p('DDA should own the central commitment and obtain named entity allocations with funded user counts, deployment type and migration dates. The initial 1,830 reported project users leave 8,170 seats awaiting allocation. Prioritize the survey’s current requirements, validate exact counts in the four largest user bands, and resolve overlaps before using those counts to fill the gap.')
p(f'At the base Year 1 cost, approximately {m["breakEven"]:,} deployed users would match the cost of buying those users at the DDA baseline, assuming 1,260 remain on-premise. The threshold includes the modeled Year 1 onboarding fee. It is a benchmark cost threshold, not proof of current cash payback. The workbook also tests slower deployment and a quarter-year of known legacy-charge overlap.')
p('Use a quarterly allocation and adoption review to track active paid users, unused capacity and upcoming renewals. Agree the rights to assign and reallocate licenses among entities in the contract. Participating-entity budget contributions can fund the annual DDA payment, but transfers within government do not create additional savings.')
h('Commercial scope and decision conditions')
p('Digital Sovereign Sign includes unlimited manual transactions, unlimited external recipients, hosting and branding. Standard 8×5 support is 10% of subscription fees; on-premise support is 20%. Onboarding costs AED 18,000 once per new entity.')
p('Automated/API transaction add-ons, Sovereign Collaboration, AI services, and Evidence based Archiving for Standalone Documents and files are separate purchases. Confirm any additional implementation, integration, migration, training and operating costs, together with tax treatment, before treating the modeled budget as an all-in commitment.')
p('The immediate decision is to sponsor central procurement and complete the allocation and commercial schedule. Confirm the added-seat price lock, contract cessation dates and prepaid-credit treatment. The model, brief and slides remain an internal proposal until those inputs are agreed.')
h('Evidence and calculation basis')
p('Commercial terms, customers and survey: user-supplied working fact register S01–S10. Entity population: DDA deparments shared service.xlsx, dda departments columns A:B only. The 91 Entities Model sheet supplies no factual inputs. Financial values: accompanying Excel model, Base case, Cost model rows 23–47. External research limits appear in the workbook Research and Entity scope sheets.')
link('Digital Dubai announcement on Smart Employee scale, May 2025', 'https://www.digitaldubai.ae/newsroom/news/dubai-to-make-presence-felt-at-gitex-europe-x-ai-everything-2025-with-a-joint-pavilion-featuring-12-government-and-private-entities')
doc.core_properties.title='DDA central procurement business case'
doc.core_properties.subject='Circularo five year central procurement proposal'
doc.core_properties.author='Circularo'
for tree in [doc._element, doc.styles.element]:
    for el in list(tree.iter(qn('w:pBdr'))):el.getparent().remove(el)
dest=root/'outputs/dda-business-case/DDA-executive-brief.docx';doc.save(dest);print(dest)
