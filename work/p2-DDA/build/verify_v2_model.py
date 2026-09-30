import json, os, subprocess, hashlib, math
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET
from openpyxl import load_workbook

root=Path(__file__).resolve().parents[1]
qa=root/'build/v2'
final=root/'outputs/dda-business-case-v2/DDA-savings-model-v2.1.xlsx'
source=root/'DDA deparments shared service.xlsx'
original_hash=hashlib.sha256(source.read_bytes()).hexdigest()
native=qa/'native-recalculated';native.mkdir(exist_ok=True)
profile=qa/'native-profile';profile.mkdir(exist_ok=True)
temp=qa/'tmp';temp.mkdir(exist_ok=True)
env=os.environ.copy()
env['TMPDIR']=str(temp)
for key,folder in [('XDG_CONFIG_HOME','config'),('XDG_CACHE_HOME','cache')]:
    d=profile/folder;d.mkdir(exist_ok=True);env[key]=str(d)
soffice='/Users/josefneumann/.cache/codex-runtimes/codex-primary-runtime/dependencies/bin/override/soffice'
files=[final,qa/'engine-input-8000.xlsx',qa/'engine-input-5000.xlsx']
command=[soffice,f'-env:UserInstallation={profile.as_uri()}','--headless','--norestore','--convert-to','xlsx','--outdir',str(native),*[str(p)for p in files]]
run=subprocess.run(command,env=env,capture_output=True,text=True,timeout=55)
print(run.stdout.strip());assert run.returncode==0,run.stderr
results=[]
for file,n,sub,support,saving in [(files[0],10000,11800000,1652000,383369.75),(files[1],8000,10620000,1699200,190144.75),(files[2],5000,7375000,1327500,-3080.25)]:
    b=load_workbook(native/file.name,data_only=True)
    errors=[(s.title,c.coordinate,c.value)for s in b for row in s for c in row if c.data_type=='e']
    assert not errors,errors
    expected={'Setup!E5':n,'Setup!E54':92,'Existing customers!D16':1310,'Existing customers!G16':1929169.75,'Existing customers!I16':saving,'Existing customers!D15':50,'Existing customers!E15':2212.5,'Existing customers!G15':110625,'Entities!C101':92,'Entities!D103':93,'Potential!D39':sub,'Potential!D41':support,'Potential!E11':15200,'Survey!H27':4921,'Survey!G27':10558.5,'Survey!G15':4550006.5,'Potential!E23':262,'Summary!D10':262,'Summary!D63':262}
    existing=1572;survey=min(4921,n-existing);other=n-existing-survey
    forecast=[]
    for i,col in enumerate('DEFGH'):
        if i:
            existing=math.floor(existing*1.05+.5);survey=math.floor(survey*1.05+.5);other=math.floor(other*1.05+.5)
        paid=max(n,existing+survey+other)
        benchmark=1929169.75*existing/1310+(survey+other)*2950
        central=paid*sub/n
        for row,value in [(28,existing),(29,survey),(30,other),(31,existing+survey+other),(32,paid),(33,sub/n),(38,benchmark),(39,central),(40,benchmark-central),(41,central*support/sub)]:
            expected[f'Potential!{col}{row}']=value
        forecast.append({'users':paid,'subscription':central,'saving':benchmark-central})
    expected['Potential!I39']=sum(x['subscription'] for x in forecast)
    expected['Potential!I40']=sum(x['saving'] for x in forecast)
    for key,value in expected.items():
        sheet,cell=key.split('!');actual=b[sheet][cell].value
        assert isinstance(actual,(float,int)) and abs(actual-value)<.01,(key,actual,value)
    assert b['Survey']['E38'].value=='n.a.'
    assert b['Summary']['E4'].value==n
    assert abs(b['Summary']['G13'].value-b['Potential']['D40'].value)<.01
    assert abs(b['Summary']['H13'].value-b['Potential']['I40'].value)<.01
    assert b['Summary']['D13'].value==n
    assert b['Summary']['F23'].value==4921
    assert b['Summary']['H23'].value==10558.5
    assert b['Survey']['E66'].value==54
    assert abs(b['Potential']['E100'].value)<.01
    results.append({'commitment':n,'checks':expected,'formula_errors':errors})

original=load_workbook(source,data_only=True,read_only=True)['dda departments']
records=[(original.cell(r,1).value,original.cell(r,2).value)for r in range(2,94)]
book=load_workbook(final,data_only=False)
assert book['Setup']['E5'].value=="='Summary'!E4"
assert isinstance(book['Summary']['E4'].value,(int,float))
before=load_workbook(qa/'before-summary-control-v2.1.xlsx',data_only=False)
allowed={('Summary',c) for c in ['C4','E4','G4']}|{('Setup',c) for c in ['C5','E5','G5','C48']}
assert before.sheetnames==book.sheetnames
for sheet in before:
    for row in sheet:
        for cell in row:
            if (sheet.title,cell.coordinate) not in allowed:
                assert cell.value==book[sheet.title][cell.coordinate].value,(sheet.title,cell.coordinate,cell.value,book[sheet.title][cell.coordinate].value)
assert records==[(book['Entities'].cell(r,3).value,book['Entities'].cell(r,4).value)for r in range(9,101)]
assert (book['Entities']['C101'].value,book['Entities']['D101'].value)==(92,'His Highness the Rulers Court')
assert book['Existing customers']['C15'].value=='His Highness the Rulers Court'
assert len(book['Summary']._charts)==2
ns={'c':'http://schemas.openxmlformats.org/drawingml/2006/chart','a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
with ZipFile(final)as z:
    charts=[n for n in z.namelist()if n.endswith('.xml')and '/charts/chart' in n]
    for name in charts:
        xml=ET.fromstring(z.read(name))
        expected_series=2 if name.endswith('chart1.xml') else 1
        assert len(xml.findall('.//c:ser',ns))==expected_series
        colors=[x.attrib['val']for x in xml.findall('.//c:ser/c:spPr/a:solidFill/a:srgbClr',ns)]
        assert colors==['1D0090','7000FF'][:expected_series],colors
        refs=[x.text for x in xml.findall('.//c:f',ns)]
        assert len(refs)>=2*expected_series,refs
        if name.endswith('chart2.xml'):
            assert any('66' in ref for ref in refs),refs
    assert len(charts)==2
assert hashlib.sha256(source.read_bytes()).hexdigest()==original_hash
(qa/'native-verification.json').write_text(json.dumps({'engine':'Bundled LibreOffice; private recalculated copies','scenarios':results,'source_rows_verified':len(records),'user_added_entities_verified':1,'native_charts_verified':2,'source_unchanged':True},indent=2))
print('10K, 8K and 5K recalculated correctly. Source rows and native charts verified.')
